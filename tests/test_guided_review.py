import hashlib
import io
import zipfile

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document as WordDocument
from docx.shared import Pt
from lxml import etree

from reviews.document_edits import NS, build_document, editing_options
from reviews.extraction import extract
from reviews.models import EditDecision, Finding
from tests.test_application import ReviewFixture


@override_settings(AI_ENABLED=False)
class GuidedReviewTests(ReviewFixture, TestCase):
    def setUp(self):
        super().setUp()
        for document in self.review.documents.all():
            document.sha256 = hashlib.sha256(document.file.read()).hexdigest()
            document.save()
        self.finding = self.ready()

    def accept(self, finding=None, target=1, text=None):
        finding = finding or self.finding
        self.review.refresh_from_db()
        option = editing_options(finding)[target]
        return self.client.post(
            reverse("review_change", args=[self.review.pk, finding.pk]),
            {
                "revision": self.review.revision,
                "choice": "accepted",
                "targets": [str(target)],
                f"after_{target}": text if text is not None else option["after"],
                "note": "User confirmed intended wording.",
            },
        )

    def decide_remaining(self):
        for finding in self.review.findings.filter(edit_decision__isnull=True):
            self.review.refresh_from_db()
            self.client.post(
                reverse("review_change", args=[self.review.pk, finding.pk]),
                {
                    "revision": self.review.revision,
                    "choice": "rejected",
                    "note": "No source change.",
                },
            )

    def test_guided_flow_accept_edit_reject_approve_download(self):
        response = self.client.get(reverse("guide", args=[self.review.pk]))
        self.assertContains(response, "Issue 1 of 3")
        self.assertContains(response, "Accept edit")
        self.assertEqual(
            self.accept(
                text="The primary endpoint is change from baseline in fatigue score at Week 12."
            ).status_code,
            302,
        )
        decision = EditDecision.objects.get(finding=self.finding)
        self.assertEqual(decision.edits[0]["document_id"], self.review.documents.get(role="sap").pk)
        self.decide_remaining()
        self.review.refresh_from_db()
        doc = self.review.documents.get(role="sap")
        url = reverse("download_document", args=[self.review.pk, doc.pk])
        self.assertEqual(self.client.get(url).status_code, 409)
        response = self.client.post(
            reverse("finalise", args=[self.review.pk]),
            {"revision": self.review.revision, "approval": "confirmed"},
        )
        self.assertEqual(response.status_code, 302)
        self.review.refresh_from_db()
        self.assertEqual(self.review.finalised_revision, self.review.revision)
        result = self.client.get(url)
        self.assertEqual(result.status_code, 200)
        exported = WordDocument(io.BytesIO(b"".join(result.streaming_content)))
        text = "\n".join(p.text for p in exported.paragraphs)
        self.assertIn("score at Week 12.", text)
        self.assertNotIn("score at Week 24.", text)
        self.assertIn("per-protocol", text)  # Rejected population suggestion remains unchanged.
        self.assertIn(
            "score at Week 24.", "\n".join(p.text for p in WordDocument(doc.file).paragraphs)
        )

    def test_word_copy_has_native_revisions_and_comment_wiring(self):
        self.accept()
        doc = self.review.documents.get(role="sap")
        stream = build_document(self.review, doc, tracked=True)
        with zipfile.ZipFile(stream) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            self.assertTrue(root.xpath("//w:del/w:r/w:delText", namespaces=NS))
            self.assertTrue(root.xpath("//w:ins/w:r/w:t", namespaces=NS))
            self.assertTrue(root.xpath("//w:commentRangeStart", namespaces=NS))
            comments = etree.fromstring(archive.read("word/comments.xml"))
            starts = set(root.xpath("//w:commentRangeStart/@w:id", namespaces=NS))
            ends = set(root.xpath("//w:commentRangeEnd/@w:id", namespaces=NS))
            references = set(root.xpath("//w:commentReference/@w:id", namespaces=NS))
            self.assertEqual(starts, ends)
            self.assertEqual(starts, references)
            self.assertEqual(starts, {c.get("{" + NS["w"] + "}id") for c in comments})
            self.assertIn(b"relationships/comments", archive.read("word/_rels/document.xml.rels"))
            self.assertIn(b"/word/comments.xml", archive.read("[Content_Types].xml"))
            self.assertIn(b"trackRevisions", archive.read("word/settings.xml"))

    def test_pending_conflicts_become_comments_without_assumed_edits(self):
        doc = self.review.documents.get(role="sap")
        stream = build_document(self.review, doc, tracked=True)
        with zipfile.ZipFile(stream) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            self.assertFalse(root.xpath("//w:ins|//w:del", namespaces=NS))
            self.assertIn(b"no edit direction has been selected", archive.read("word/comments.xml"))

    def test_stale_decision_and_final_approval_are_rejected(self):
        self.accept()
        response = self.client.post(
            reverse("review_change", args=[self.review.pk, self.finding.pk]),
            {"revision": 0, "choice": "rejected"},
        )
        self.assertEqual(response.status_code, 409)
        self.assertEqual(
            self.client.post(
                reverse("finalise", args=[self.review.pk]), {"revision": 0, "approval": "confirmed"}
            ).status_code,
            409,
        )

    def test_all_findings_and_explicit_approval_required(self):
        self.review.refresh_from_db()
        url = reverse("finalise", args=[self.review.pk])
        self.assertEqual(
            self.client.post(
                url, {"revision": self.review.revision, "approval": "confirmed"}
            ).status_code,
            409,
        )
        self.decide_remaining()
        self.review.refresh_from_db()
        self.assertEqual(self.client.post(url, {"revision": self.review.revision}).status_code, 400)

    def test_revising_a_decision_invalidates_final_download(self):
        self.accept()
        self.decide_remaining()
        self.review.refresh_from_db()
        self.client.post(
            reverse("finalise", args=[self.review.pk]),
            {"revision": self.review.revision, "approval": "confirmed"},
        )
        self.accept(text="The primary endpoint is the fatigue score change at Week 12.")
        doc = self.review.documents.get(role="sap")
        self.assertEqual(
            self.client.get(
                reverse("download_document", args=[self.review.pk, doc.pk])
            ).status_code,
            409,
        )

    def test_overlap_rejected_without_losing_submitted_draft(self):
        self.accept()
        duplicate = Finding.objects.create(
            review=self.review,
            check_key="analysis",
            kind="question",
            priority="medium",
            title="Related wording",
            why="Same source",
            action="Confirm",
            unknown="Intent",
            evidence=self.finding.evidence,
        )
        response = self.accept(finding=duplicate, text="Keep this draft wording.")
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "Keep this draft wording.", status_code=400)
        self.assertFalse(EditDecision.objects.filter(finding=duplicate).exists())
        self.assertFalse(any("edit_draft" in key for key in self.client.session.keys()))

    def test_cross_project_edit_export_and_approval_denied(self):
        self.client.force_login(self.outsider)
        doc = self.review.documents.get(role="sap")
        for url in [
            reverse("guide", args=[self.review.pk]),
            reverse("download_document", args=[self.review.pk, doc.pk]),
        ]:
            self.assertEqual(self.client.get(url).status_code, 404)
        for url in [
            reverse("finalise", args=[self.review.pk]),
            reverse("review_change", args=[self.review.pk, self.finding.pk]),
        ]:
            self.assertEqual(self.client.post(url).status_code, 404)

    def test_no_selection_is_actionable_and_preserves_source(self):
        response = self.client.post(
            reverse("review_change", args=[self.review.pk, self.finding.pk]),
            {
                "revision": self.review.revision,
                "choice": "accepted",
                "after_1": "My replacement draft",
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertContains(response, "My replacement draft", status_code=400)
        self.assertFalse(EditDecision.objects.exists())

    def test_original_hash_tamper_prevents_edit(self):
        document = self.review.documents.get(role="sap")
        document.sha256 = "0" * 64
        document.save()
        self.assertFalse(editing_options(self.finding)[1]["editable"])

    def test_question_can_be_kept_as_native_comment(self):
        response = self.client.post(
            reverse("review_change", args=[self.review.pk, self.finding.pk]),
            {
                "revision": self.review.revision,
                "choice": "commented",
                "note": "Confirm the primary time with the statistician.",
            },
        )
        self.assertEqual(response.status_code, 302)
        stream = build_document(self.review, self.review.documents.get(role="sap"))
        with zipfile.ZipFile(stream) as archive:
            self.assertIn(b"Confirm the primary time", archive.read("word/comments.xml"))
            self.assertNotIn(b"<w:ins ", archive.read("word/document.xml"))

    def test_unchanged_replacement_not_silently_accepted(self):
        self.assertEqual(self.accept(text=self.finding.evidence[1]["quote"]).status_code, 400)

    def test_formatted_runs_and_unmodified_package_parts_survive(self):
        doc = WordDocument()
        p = doc.add_paragraph()
        p.add_run("Before ").bold = True
        p.add_run("Week ").italic = True
        p.add_run("24").font.size = Pt(14)
        p.add_run(" after.").underline = True
        doc.sections[0].header.paragraphs[0].text = "Preserved header"
        stream = io.BytesIO()
        doc.save(stream)
        content = stream.getvalue()
        result = extract(content, "format.docx")
        # Reuse one role while retaining the exact immutable source bytes for this fixture.
        document = self.review.documents.get(role="protocol")
        document.file.save("format.docx", ContentFile(content), save=False)
        document.sha256 = hashlib.sha256(content).hexdigest()
        document.blocks = result["blocks"]
        document.save()
        evidence = {
            "block_id": f"d{document.pk}:b1",
            "document_id": document.pk,
            "name": document.name,
            "version": "1",
            "location": "DOCX body / block 1",
            "start": 7,
            "end": 14,
            "quote": "Week 24",
            "after": "Week 12",
        }
        finding = Finding.objects.create(
            review=self.review,
            check_key="endpoint",
            kind="question",
            priority="medium",
            title="Formatting",
            why="Timing",
            action="Confirm",
            unknown="Time",
            evidence=[{k: v for k, v in evidence.items() if k != "after"}],
        )
        EditDecision.objects.create(
            finding=finding, choice="accepted", edits=[evidence], actor=self.user
        )
        exported = build_document(self.review, document)
        result_doc = WordDocument(exported)
        self.assertEqual(result_doc.paragraphs[0].text, "Before Week 12 after.")
        self.assertTrue(result_doc.paragraphs[0].runs[0].bold)
        self.assertTrue(result_doc.paragraphs[0].runs[-1].underline)
        with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(exported) as updated:
            self.assertEqual(original.read("word/header1.xml"), updated.read("word/header1.xml"))

    def test_invalid_guided_step_is_not_server_error(self):
        self.assertEqual(
            self.client.get(reverse("guide", args=[self.review.pk]) + "?step=oops").status_code, 400
        )

    def test_selecting_both_conflicting_directions_is_rejected(self):
        options = editing_options(self.finding)
        response = self.client.post(
            reverse("review_change", args=[self.review.pk, self.finding.pk]),
            {
                "revision": self.review.revision,
                "choice": "accepted",
                "targets": ["0", "1"],
                "after_0": options[0]["after"],
                "after_1": options[1]["after"],
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(EditDecision.objects.exists())

    def test_native_revision_accept_and_reject_reconstruct_expected_text(self):
        self.accept()
        doc = self.review.documents.get(role="sap")
        content = build_document(self.review, doc, tracked=True).getvalue()
        for accept, expected in [(True, "Week 12"), (False, "Week 24")]:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                root = etree.fromstring(archive.read("word/document.xml"))
                for node in list(root.xpath("//w:ins|//w:del", namespaces=NS)):
                    parent = node.getparent()
                    if (node.tag.endswith("ins") and accept) or (
                        node.tag.endswith("del") and not accept
                    ):
                        at = parent.index(node)
                        for child in list(node):
                            for text in child.findall(".//{" + NS["w"] + "}delText"):
                                text.tag = "{" + NS["w"] + "}t"
                            parent.insert(at, child)
                            at += 1
                    parent.remove(node)
                paragraphs = root.xpath("//w:body/w:p", namespaces=NS)
                endpoint = next(
                    p for p in paragraphs if "The primary endpoint is" in "".join(p.itertext())
                )
                self.assertIn(expected, "".join(endpoint.itertext()))

    def test_table_question_anchors_a_comment_without_rewriting_cells(self):
        doc = self.review.documents.get(role="protocol")
        block = next(b for b in doc.blocks if b["kind"] == "table")
        evidence = {
            "block_id": f"d{doc.pk}:{block['id']}",
            "document_id": doc.pk,
            "name": doc.name,
            "version": "1",
            "location": block["location"],
            "start": 0,
            "end": len(block["text"]),
            "quote": block["text"],
        }
        finding = Finding.objects.create(
            review=self.review,
            check_key="schedule",
            kind="question",
            priority="medium",
            title="Confirm schedule",
            why="Confirm timing",
            action="Check the visit.",
            unknown="Visit",
            evidence=[evidence],
        )
        self.assertFalse(editing_options(finding)[0]["editable"])
        EditDecision.objects.create(
            finding=finding, choice="commented", note="Check this table.", actor=self.user
        )
        content = build_document(self.review, doc).getvalue()
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            xml = etree.fromstring(archive.read("word/document.xml"))
            self.assertTrue(xml.xpath("//w:tbl//w:commentReference", namespaces=NS))
        self.assertEqual(WordDocument(io.BytesIO(content)).tables[0].cell(0, 1).text, "Baseline")

    def test_existing_comments_and_untouched_formatting_are_preserved(self):
        self.review.findings.all().delete()
        source = WordDocument()
        old = source.add_paragraph("Keep the existing reviewer comment.")
        source.add_comment(old.runs, text="Existing reviewer question", author="Earlier Reviewer")
        source.add_paragraph("Change this wording.")
        stream = io.BytesIO()
        source.save(stream)
        content = stream.getvalue()
        doc = self.review.documents.get(role="protocol")
        doc.file.save("comments.docx", ContentFile(content), save=False)
        doc.sha256 = hashlib.sha256(content).hexdigest()
        doc.blocks = extract(content, "comments.docx")["blocks"]
        doc.save()
        block = doc.blocks[1]
        evidence = {
            "block_id": f"d{doc.pk}:{block['id']}",
            "document_id": doc.pk,
            "name": doc.name,
            "version": "1",
            "location": block["location"],
            "start": 0,
            "end": len(block["text"]),
            "quote": block["text"],
        }
        finding = Finding.objects.create(
            review=self.review,
            check_key="operations",
            kind="question",
            priority="medium",
            title="New wording",
            why="Clarify",
            action="Confirm",
            unknown="Wording",
            evidence=[evidence],
        )
        EditDecision.objects.create(
            finding=finding,
            choice="accepted",
            edits=[{**evidence, "after": "Updated wording."}],
            actor=self.user,
        )
        content = build_document(self.review, doc, tracked=True).getvalue()
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            comments = etree.fromstring(archive.read("word/comments.xml"))
            self.assertIn("Existing reviewer question", "".join(comments.itertext()))
            self.assertEqual(len(comments), 2)
            ids = [c.get("{" + NS["w"] + "}id") for c in comments]
            self.assertEqual(len(ids), len(set(ids)))

    def test_quick_upload_needs_no_project_or_review_name(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        from reviews.management.commands.seed_demo import demo_doc
        from reviews.models import Review

        response = self.client.post(
            reverse("quick_upload"),
            {
                "protocol": SimpleUploadedFile("my_trial.docx", demo_doc("protocol")),
                "authorised": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        new = Review.objects.get(title="my trial")
        self.assertIn("/guide/", response["Location"])
        self.assertTrue(new.project.members.filter(pk=self.user.pk).exists())

    def test_pdf_cannot_be_downloaded_as_a_rebuilt_word_original(self):
        doc = self.review.documents.get(role="sap")
        doc.name = "original.pdf"
        doc.save()
        response = self.client.get(
            reverse("download_document", args=[self.review.pk, doc.pk]) + "?mode=review"
        )
        self.assertEqual(response.status_code, 400)

    def test_word_handoff_can_propose_pending_edits_with_explicit_reference(self):
        doc = self.review.documents.get(role="sap")
        content = build_document(self.review, doc, tracked=True, reference="protocol").getvalue()
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            self.assertEqual(len(root.xpath("//w:ins", namespaces=NS)), 3)
            self.assertIn(b"no decision has been saved", archive.read("word/comments.xml"))
        self.assertFalse(EditDecision.objects.exists())
        self.assertEqual(self.review.revision, 0)

    def test_word_handoff_respects_rejected_decisions(self):
        self.decide_remaining()
        doc = self.review.documents.get(role="sap")
        with zipfile.ZipFile(
            build_document(self.review, doc, tracked=True, reference="protocol")
        ) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            self.assertFalse(root.xpath("//w:ins|//w:del", namespaces=NS))

    def test_project_deletion_cascades_to_edit_decisions(self):
        self.accept()
        self.assertTrue(EditDecision.objects.exists())
        self.client.post(
            reverse("delete_project", args=[self.project.pk]), {"confirm": str(self.project.pk)}
        )
        self.assertFalse(EditDecision.objects.exists())

    def test_a_workspace_record_change_invalidates_final_approval(self):
        self.decide_remaining()
        self.review.refresh_from_db()
        self.client.post(
            reverse("finalise", args=[self.review.pk]),
            {"revision": self.review.revision, "approval": "confirmed"},
        )
        finding = self.review.findings.first()
        self.client.post(
            reverse("decision", args=[self.review.pk, finding.pk]),
            {
                "revision": finding.revision,
                "status": "needs_discussion",
                "note": "A further discussion is needed.",
            },
        )
        doc = self.review.documents.get(role="sap")
        self.assertEqual(
            self.client.get(
                reverse("download_document", args=[self.review.pk, doc.pk])
            ).status_code,
            409,
        )

    def test_reference_based_word_download_is_available_before_approval(self):
        doc = self.review.documents.get(role="sap")
        url = reverse("download_document", args=[self.review.pk, doc.pk])
        response = self.client.get(url + "?mode=review&reference=protocol")
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(b"".join(response.streaming_content))) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
            self.assertEqual(len(root.xpath("//w:ins", namespaces=NS)), 3)
        self.assertFalse(EditDecision.objects.exists())

    def test_unknown_word_reference_is_rejected(self):
        doc = self.review.documents.get(role="sap")
        url = reverse("download_document", args=[self.review.pk, doc.pk])
        self.assertEqual(self.client.get(url + "?mode=review&reference=elsewhere").status_code, 400)
