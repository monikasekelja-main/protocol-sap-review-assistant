import io
import json
import tempfile
import zipfile
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document as WordDocument

from reviews.extraction import ExtractionError, extract
from reviews.jobs import process_next
from reviews.management.commands.seed_demo import demo_doc
from reviews.models import Document, Event, Finding, Message, Project, Proposal, Review
from reviews.pipeline import local_review, run_review, scope_signals, source_blocks
from reviews.provider import ProviderError, generate, provider_schema
from reviews.schemas import Candidate, DiscussionOutput, ReviewOutput, validate_finding


class ExtractionTests(SimpleTestCase):
    def test_docx_keeps_locations_and_table_rows(self):
        result = extract(demo_doc("protocol"), "protocol.docx")
        self.assertTrue(any("Week 12" in b["text"] for b in result["blocks"]))
        rows = [b for b in result["blocks"] if b["kind"] == "table"]
        self.assertEqual(rows[0]["text"], "Assessment | Baseline | Week 12")
        self.assertIn("table row 1", rows[0]["location"])
        self.assertTrue(any("merged headers" in reason for reason in result["limitations"]))
        self.assertEqual(result["metadata"]["version"], "1.0")

    def test_fake_extension_and_empty_rejected(self):
        for content, name in [(b"not a pdf", "test.pdf"), (b"PKfake", "test.txt")]:
            with self.assertRaises(ExtractionError):
                extract(content, name)
        stream = io.BytesIO()
        WordDocument().save(stream)
        with self.assertRaises(ExtractionError):
            extract(stream.getvalue(), "empty.docx")

    def test_tracked_changes_rejected(self):
        source = zipfile.ZipFile(io.BytesIO(demo_doc("protocol")))
        target = io.BytesIO()
        with zipfile.ZipFile(target, "w") as archive:
            for entry in source.infolist():
                content = source.read(entry.filename)
                if entry.filename == "word/document.xml":
                    content = content.replace(b"<w:body>", b"<w:body><w:ins/>")
                archive.writestr(entry, content)
        with self.assertRaisesMessage(ExtractionError, "tracked changes"):
            extract(target.getvalue(), "tracked.docx")

    def test_oversized_archive_rejected(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"x" * (101 * 1024 * 1024))
        with self.assertRaisesMessage(ExtractionError, "expansion"):
            extract(stream.getvalue(), "large.docx")


class ReviewFixture:
    def setUp(self):
        self.storage = tempfile.TemporaryDirectory()
        self.addCleanup(self.storage.cleanup)
        self.override = override_settings(MEDIA_ROOT=self.storage.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = get_user_model().objects.create_user("reviewer", password="test-only-pass")
        self.outsider = get_user_model().objects.create_user("outsider")
        self.project = Project.objects.create(name="Study", owner=self.user)
        self.project.members.add(self.user)
        self.review = Review.objects.create(
            project=self.project, title="Test review", scope_confirmed=True, metadata_confirmed=True
        )
        for role in ("protocol", "sap"):
            content = demo_doc(role)
            result = extract(content, role + ".docx")
            Document.objects.create(
                review=self.review,
                role=role,
                name=role + ".docx",
                file=ContentFile(content, name=role + ".docx"),
                sha256="a" * 64,
                **result,
            )
        self.client.force_login(self.user)

    def ready(self):
        self.assertTrue(process_next())
        self.review.refresh_from_db()
        self.assertEqual(self.review.state, "ready", self.review.error)
        return self.review.findings.first()


@override_settings(AI_ENABLED=False)
class ApplicationTests(ReviewFixture, TestCase):
    def test_end_to_end_review_decision_and_export(self):
        finding = self.ready()
        self.assertEqual(self.review.findings.count(), 3)
        self.assertEqual(len(self.review.coverage), 8)
        response = self.client.get(reverse("review", args=[self.review.pk]))
        self.assertContains(response, "Review ready")
        self.assertContains(response, "Primary assessment time differs")
        url = reverse("decision", args=[self.review.pk, finding.pk])
        self.assertEqual(
            self.client.post(
                url, {"status": "action_needed", "note": "Week 12 confirmed", "revision": 0}
            ).status_code,
            302,
        )
        self.assertEqual(
            self.client.post(
                url, {"status": "dismissed", "note": "stale", "revision": 0}
            ).status_code,
            409,
        )
        finding.refresh_from_db()
        self.assertEqual(finding.note, "Week 12 confirmed")
        self.assertEqual(Event.objects.filter(kind="decision").count(), 1)
        report = self.client.get(reverse("report", args=[self.review.pk]))
        doc = WordDocument(io.BytesIO(b"".join(report.streaming_content)))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Week 12 confirmed", text)
        self.assertIn("Coverage and limitations", text)
        self.assertIn("SHA-256", text)
        self.assertNotIn("stale", text)

    def test_cross_project_access_is_denied(self):
        finding = self.ready()
        doc = self.review.documents.first()
        self.client.force_login(self.outsider)
        urls = [
            reverse("review", args=[self.review.pk]),
            reverse("document", args=[self.review.pk, doc.pk]),
            reverse("report", args=[self.review.pk]),
            reverse("upload", args=[self.project.pk]),
        ]
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 404)
        for url in [
            reverse("decision", args=[self.review.pk, finding.pk]),
            reverse("chat", args=[self.review.pk]),
            reverse("propose", args=[self.review.pk]),
            reverse("delete_project", args=[self.project.pk]),
            reverse("add_member", args=[self.project.pk]),
        ]:
            self.assertEqual(self.client.post(url).status_code in (400, 404), True)

    def test_anonymous_redirects(self):
        self.client.logout()
        self.assertEqual(self.client.get("/").status_code, 302)

    def test_document_download_is_private_and_unchanged(self):
        doc = self.review.documents.first()
        response = self.client.get(
            reverse("document", args=[self.review.pk, doc.pk]) + "?download=1"
        )
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(b"".join(response.streaming_content), doc.file.read())
        self.assertEqual(self.client.get("/media/" + doc.file.name).status_code, 404)

    def test_clarification_requires_acceptance_and_does_not_resolve_finding(self):
        finding = self.ready()
        self.client.post(
            reverse("propose", args=[self.review.pk]),
            {"finding": finding.pk, "text": "We intend Week 12."},
        )
        proposal = Proposal.objects.get()
        self.assertFalse(Event.objects.filter(kind="clarification").exists())
        response = self.client.post(
            reverse("accept_proposal", args=[self.review.pk, proposal.pk]), {"choice": "accept"}
        )
        self.assertEqual(response.status_code, 302)
        finding.refresh_from_db()
        self.assertEqual(finding.status, "not_reviewed")
        self.assertTrue(Event.objects.filter(kind="clarification").exists())
        self.assertEqual(
            self.client.post(
                reverse("accept_proposal", args=[self.review.pk, proposal.pk]), {"choice": "accept"}
            ).status_code,
            409,
        )

    def test_stale_proposal_rejected_and_rejection_no_mutation(self):
        finding = self.ready()
        proposal = Proposal.objects.create(
            review=self.review,
            finding=finding,
            kind="clarification",
            payload={"note": "Context"},
            base_revision=0,
        )
        finding.revision = 1
        finding.save()
        url = reverse("accept_proposal", args=[self.review.pk, proposal.pk])
        self.assertEqual(self.client.post(url, {"choice": "accept"}).status_code, 409)
        self.assertEqual(self.client.post(url, {"choice": "reject"}).status_code, 302)
        self.assertFalse(Event.objects.filter(kind="clarification").exists())

    def test_local_discussion_is_grounded_and_never_changes_records(self):
        finding = self.ready()
        response = self.client.post(
            reverse("chat", args=[self.review.pk]),
            {"finding": finding.pk, "message": "Why did you flag Week 12?"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Message.objects.count(), 2)
        self.assertTrue(Message.objects.get(role="assistant").evidence)
        self.assertEqual(Proposal.objects.count(), 0)
        finding.refresh_from_db()
        self.assertEqual(finding.revision, 0)

    def test_general_chat_has_separate_context(self):
        self.ready()
        self.client.post(reverse("chat", args=[self.review.pk]), {"message": "Show the population"})
        self.assertEqual(Message.objects.filter(finding__isnull=True).count(), 2)
        response = self.client.get(reverse("review", args=[self.review.pk]) + "?general=1")
        self.assertContains(response, "Show the population")
        self.assertContains(response, "Ask about your documents")

    def test_delete_removes_files_and_all_records(self):
        self.ready()
        paths = [Path(d.file.path) for d in self.review.documents.all()]
        self.assertTrue(all(p.exists() for p in paths))
        self.assertEqual(
            self.client.post(
                reverse("delete_project", args=[self.project.pk]), {"confirm": str(self.project.pk)}
            ).status_code,
            302,
        )
        self.assertFalse(any(p.exists() for p in paths))
        for model in (Project, Review, Document, Finding, Event, Message, Proposal):
            self.assertEqual(model.objects.count(), 0)

    def test_delete_failure_revokes_access_and_can_retry(self):
        with patch("django.db.models.fields.files.FieldFile.delete", side_effect=OSError):
            self.client.post(
                reverse("delete_project", args=[self.project.pk]), {"confirm": str(self.project.pk)}
            )
        self.assertEqual(self.client.get(reverse("review", args=[self.review.pk])).status_code, 404)
        self.client.post(
            reverse("delete_project", args=[self.project.pk]), {"confirm": str(self.project.pk)}
        )
        self.assertFalse(Project.objects.exists())

    def test_worker_extracts_in_bounded_subprocess_and_asks_for_confirmation(self):
        self.review.scope_confirmed = self.review.metadata_confirmed = False
        self.review.save()
        self.review.documents.update(blocks=[], metadata={}, limitations=[])
        self.assertTrue(process_next())
        self.review.refresh_from_db()
        self.assertEqual(self.review.state, "needs_input", self.review.error)
        self.assertTrue(self.review.documents.first().blocks)
        response = self.client.get(reverse("review", args=[self.review.pk]))
        self.assertContains(response, "Are these the right versions?")
        self.assertEqual(
            self.client.post(
                reverse("confirm", args=[self.review.pk]),
                {"scope": "confirmed", "metadata": "confirmed"},
            ).status_code,
            302,
        )
        self.assertEqual(self.review.findings.count(), 0)
        self.ready()

    def test_unusable_file_fails_without_erasing_other_runs(self):
        self.ready()
        other = Review.objects.create(project=self.project, title="Bad document")
        Document.objects.create(
            review=other,
            role="protocol",
            name="bad.pdf",
            file=ContentFile(b"not pdf", name="bad.pdf"),
            sha256="b" * 64,
        )
        process_next()
        other.refresh_from_db()
        self.assertEqual(other.state, "failed")
        self.assertEqual(self.review.findings.count(), 3)

    def test_timed_out_job_preserves_decisions(self):
        finding = self.ready()
        self.review.state = "processing"
        self.review.lease_until = timezone.now() - timedelta(seconds=1)
        self.review.save()
        process_next()
        self.review.refresh_from_db()
        self.assertEqual(self.review.state, "failed")
        self.assertTrue(Finding.objects.filter(pk=finding.pk).exists())

    def test_stale_worker_cannot_publish(self):
        def steal_lease(review):
            Review.objects.filter(pk=review.pk).update(lease_token=None, state="cancelled")
            return {"findings": [], "summary": {}, "coverage": [], "limitations": [], "usage": {}}

        with patch("reviews.jobs.run_review", side_effect=steal_lease):
            process_next()
        self.review.refresh_from_db()
        self.assertEqual(self.review.state, "cancelled")
        self.assertFalse(self.review.findings.exists())

    def test_wrong_version_and_single_passage_discrepancies_rejected(self):
        blocks = source_blocks(self.review)
        b = next(iter(blocks.values()))
        candidate = Candidate(
            check="endpoint",
            kind="discrepancy",
            priority="high",
            title="Mismatch",
            why="Consequence",
            action="Confirm",
            unknown="",
            evidence=[{"block_id": b["id"], "quote": b["text"]}],
        )
        with self.assertRaises(ValueError):
            validate_finding(candidate, blocks)
        candidate.evidence[0].block_id = "other-version:b1"
        with self.assertRaises(ValueError):
            validate_finding(candidate, blocks)

    def test_equivalent_times_do_not_trigger_discrepancy(self):
        for doc in self.review.documents.all():
            for block in doc.blocks:
                block["text"] = block["text"].replace("Week 24", "Day 84")
            doc.save()
        _, candidates, _ = local_review(source_blocks(self.review))
        self.assertFalse(any(c.check == "endpoint" for c in candidates))

    def test_secondary_endpoints_do_not_trigger_primary_comparison(self):
        self.review.documents.filter(role="sap").update(
            blocks=[
                {
                    "id": "b1",
                    "text": "The secondary endpoint is change from baseline in fatigue score at Week 24.",
                    "location": "body",
                    "kind": "paragraph",
                }
            ]
        )
        _, candidates, _ = local_review(source_blocks(self.review))
        self.assertFalse(any(c.check == "endpoint" for c in candidates))

    def test_scope_blocks_known_unsupported_design(self):
        doc = self.review.documents.get(role="protocol")
        doc.blocks.append(
            {
                "id": "extra",
                "text": "This is a crossover trial.",
                "location": "body",
                "kind": "paragraph",
            }
        )
        doc.save()
        self.assertTrue(scope_signals(source_blocks(self.review))[0])
        process_next()
        self.review.refresh_from_db()
        self.assertEqual(self.review.state, "unsupported")

    def test_protocol_only_discloses_scope_and_incomplete_checks(self):
        self.review.documents.filter(role="sap").delete()
        result = run_review(self.review)
        self.assertTrue(any("Protocol-only" in x for x in result["limitations"]))
        self.assertEqual(len(result["coverage"]), 8)
        self.assertTrue(any(x["state"] == "not_assessed" for x in result["coverage"]))

    def test_upload_rejects_unsupported_input_and_missing_attestation(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        response = self.client.post(
            reverse("upload", args=[self.project.pk]),
            {"title": "Wrong", "protocol": SimpleUploadedFile("scan.png", b"bad")},
        )
        self.assertContains(response, "Use a DOCX or text PDF")
        self.assertEqual(Review.objects.count(), 1)

    @override_settings(AI_ENABLED=True)
    def test_model_candidate_with_fabricated_quote_is_withheld(self):
        candidate = Candidate(
            check="endpoint",
            kind="question",
            priority="high",
            title="Unsupported",
            why="A consequence",
            action="Confirm timing",
            unknown="Which time?",
            evidence=[{"block_id": "invented", "quote": "invented"}],
        )
        result = ReviewOutput(summary=[], findings=[candidate], coverage=[])
        with patch("reviews.pipeline.generate", return_value=(result, {})):
            output = run_review(self.review)
        self.assertFalse(any(f["title"] == "Unsupported" for f in output["findings"]))
        self.assertTrue(any("withheld" in x for x in output["limitations"]))

    def test_csrf_and_unsafe_get_protected(self):
        from django.test import Client

        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(
            client.post(reverse("create_project"), {"name": "Forbidden"}).status_code, 403
        )
        self.assertEqual(self.client.get(reverse("create_project")).status_code, 405)


class ProviderTests(SimpleTestCase):
    def test_schema_has_no_unsupported_size_bounds(self):
        text = json.dumps(provider_schema(DiscussionOutput))
        self.assertNotIn("maxLength", text)
        self.assertNotIn("maxItems", text)
        self.assertIn("additionalProperties", text)

    @override_settings(AI_ENABLED=True, ANTHROPIC_API_KEY="test", ANTHROPIC_MODEL="test")
    def test_incomplete_provider_response_rejected(self):
        response = io.BytesIO(json.dumps({"stop_reason": "max_tokens", "content": []}).encode())
        with patch("urllib.request.urlopen", return_value=response):
            with self.assertRaises(ProviderError):
                generate("system", {}, DiscussionOutput)


def minimal_pdf(text):
    content = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(output)
    output += b"xref\n0 6\n0000000000 65535 f \n"
    output += b"".join(f"{n:010d} 00000 n \n".encode() for n in offsets[1:])
    output += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return output


class PdfTests(SimpleTestCase):
    def test_text_pdf_preserves_page_and_coordinates(self):
        result = extract(minimal_pdf("Primary endpoint at Week 12"), "study.pdf")
        block = result["blocks"][0]
        self.assertEqual(block["text"], "Primary endpoint at Week 12")
        self.assertIn("PDF page 1", block["location"])
        self.assertIn("bbox", block["location"])
        self.assertTrue(result["limitations"])

    def test_image_only_pdf_rejected(self):
        with self.assertRaisesMessage(ExtractionError, "No usable text"):
            extract(minimal_pdf(""), "scan.pdf")


@override_settings(AI_ENABLED=False)
class RecordIntegrityTests(ReviewFixture, TestCase):
    def test_rerun_preserves_previous_decision(self):
        finding = self.ready()
        finding.status, finding.note = "action_needed", "Keep this decision"
        finding.save()
        response = self.client.post(reverse("rerun", args=[self.review.pk]))
        self.assertEqual(response.status_code, 302)
        new = Review.objects.exclude(pk=self.review.pk).get()
        process_next()
        new.refresh_from_db()
        self.assertEqual(new.state, "ready", new.error)
        finding.refresh_from_db()
        self.assertEqual(finding.note, "Keep this decision")
        self.assertTrue(all(f.status == "not_reviewed" and not f.note for f in new.findings.all()))
        self.assertEqual(new.documents.first().sha256, self.review.documents.first().sha256)

    def test_malformed_finding_ids_are_404(self):
        self.ready()
        self.assertEqual(
            self.client.get(reverse("review", args=[self.review.pk]) + "?finding=abc").status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                reverse("chat", args=[self.review.pk]), {"finding": "abc", "message": "Why?"}
            ).status_code,
            404,
        )

    def test_model_decision_is_pending_until_accepted(self):
        finding = self.ready()
        proposed = {
            "kind": "decision",
            "status": "action_needed",
            "note": "Confirm Week 12.",
            "finding": None,
        }
        with patch("reviews.views.discuss", return_value=("Proposed decision.", [], proposed)):
            self.client.post(
                reverse("chat", args=[self.review.pk]),
                {"finding": finding.pk, "message": "Record action needed"},
            )
        finding.refresh_from_db()
        self.assertEqual(finding.status, "not_reviewed")
        proposal = Proposal.objects.get()
        self.client.post(
            reverse("accept_proposal", args=[self.review.pk, proposal.pk]), {"choice": "accept"}
        )
        finding.refresh_from_db()
        self.assertEqual(finding.status, "action_needed")
        self.assertEqual(finding.note, "Confirm Week 12.")

    def test_new_finding_requires_valid_evidence_and_explicit_acceptance(self):
        self.ready()
        source = next(iter(source_blocks(self.review).values()))
        candidate = Candidate(
            check="operations",
            kind="question",
            priority="medium",
            title="Confirm coordinator coverage",
            why="The stated collection depends on coordinator availability.",
            action="Confirm who covers coordinator absence.",
            unknown="Who covers absence?",
            evidence=[{"block_id": source["id"], "quote": source["text"]}],
        )
        proposal = Proposal.objects.create(
            review=self.review,
            kind="new_finding",
            base_revision=0,
            payload={
                "finding": validate_finding(candidate, source_blocks(self.review)),
                "note": "New question",
            },
        )
        before = self.review.findings.count()
        self.client.post(
            reverse("accept_proposal", args=[self.review.pk, proposal.pk]), {"choice": "accept"}
        )
        self.assertEqual(self.review.findings.count(), before + 1)
        self.assertTrue(self.review.findings.filter(title="Confirm coordinator coverage").exists())

    def test_user_content_is_escaped(self):
        finding = self.ready()
        Message.objects.create(
            review=self.review, finding=finding, role="user", text='<script>alert("x")</script>'
        )
        response = self.client.get(reverse("review", args=[self.review.pk]))
        self.assertNotContains(response, "<script>")
        self.assertContains(response, "&lt;script&gt;")

    def test_summary_contains_endpoint_definition_and_discloses_conflict(self):
        self.ready()
        summary = self.review.summary
        self.assertIn("fatigue score", summary["Primary outcome"]["text"])
        self.assertIn("Conflicting", summary["Primary outcome"]["text"])
        self.assertIn("Week 12", summary["Primary outcome"]["text"])
        self.assertIn("Week 24", summary["Primary outcome"]["text"])
        self.assertIn("analysis of covariance", summary["Primary analysis"]["text"])


@override_settings(AI_ENABLED=False)
class AccessAndUploadTests(ReviewFixture, TestCase):
    def test_successful_upload_creates_private_queued_run(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        response = self.client.post(
            reverse("upload", args=[self.project.pk]),
            {
                "title": "Fresh upload",
                "protocol": SimpleUploadedFile("protocol.docx", demo_doc("protocol")),
                "authorised": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        fresh = Review.objects.get(title="Fresh upload")
        self.assertEqual(fresh.state, "queued")
        self.assertEqual(fresh.documents.count(), 1)
        self.assertFalse(fresh.scope_confirmed)
        document = fresh.documents.get()
        import hashlib

        self.assertEqual(document.sha256, hashlib.sha256(document.file.read()).hexdigest())
        self.assertTrue(Path(document.file.path).exists())

    def test_owner_can_grant_and_revoke_project_access(self):
        self.client.post(
            reverse("add_member", args=[self.project.pk]), {"username": self.outsider.username}
        )
        self.assertTrue(self.project.members.filter(pk=self.outsider.pk).exists())
        self.client.post(
            reverse("remove_member", args=[self.project.pk]), {"user_id": self.outsider.pk}
        )
        self.assertFalse(self.project.members.filter(pk=self.outsider.pk).exists())
        self.assertEqual(
            self.client.post(
                reverse("remove_member", args=[self.project.pk]), {"user_id": self.user.pk}
            ).status_code,
            400,
        )
        self.assertTrue(self.project.members.filter(pk=self.user.pk).exists())

    def test_non_owner_cannot_change_access(self):
        self.project.members.add(self.outsider)
        self.client.force_login(self.outsider)
        self.assertEqual(
            self.client.post(
                reverse("remove_member", args=[self.project.pk]), {"user_id": self.user.pk}
            ).status_code,
            404,
        )
        self.assertEqual(
            self.client.post(
                reverse("add_member", args=[self.project.pk]), {"username": self.user.username}
            ).status_code,
            404,
        )


class UploadBoundaryTests(SimpleTestCase):
    def test_upload_handler_stops_at_file_limit_before_disk_write(self):
        from django.core.files.uploadhandler import StopUpload

        from reviews.upload_handlers import BoundedUploadHandler

        handler = BoundedUploadHandler()
        handler.new_file("protocol", "protocol.docx", "application/octet-stream", None)
        self.assertEqual(handler.receive_data_chunk(b"x", 0), b"x")
        handler.file_bytes = 25 * 1024 * 1024
        with self.assertRaises(StopUpload):
            handler.receive_data_chunk(b"x", 1)

    def test_oversized_content_length_is_rejected(self):
        from django.core.exceptions import RequestDataTooBig

        from reviews.upload_handlers import BoundedUploadHandler

        with self.assertRaises(RequestDataTooBig):
            BoundedUploadHandler().handle_raw_input(None, {}, 53 * 1024 * 1024, b"boundary")
