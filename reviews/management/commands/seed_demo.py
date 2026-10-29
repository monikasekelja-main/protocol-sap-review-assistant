import hashlib
import io

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.utils.crypto import get_random_string
from docx import Document as WordDocument

from reviews.models import Document, Project, Review


def demo_doc(role):
    doc = WordDocument()
    doc.add_heading("RESTORE · Synthetic trial " + role, 0)
    doc.add_paragraph("Version 1.0 · 2026-10-29")
    doc.add_paragraph(
        "Objective: compare the change in fatigue score between treatment and placebo in adults with a rare condition."
    )
    doc.add_paragraph(
        "This English-language study is a two-arm, parallel-group randomised superiority trial with a continuous fixed-time primary endpoint."
    )
    doc.add_paragraph(
        "Eligible population: adults aged 18 years and above with the study condition."
    )
    doc.add_paragraph(
        "Treatment comparison: investigational treatment versus placebo, allocated 1:1."
    )
    doc.add_paragraph(
        f"The primary endpoint is change from baseline in fatigue score at Week {12 if role == 'protocol' else 24}."
    )
    doc.add_paragraph(
        "The primary analysis population is "
        + ("intention-to-treat." if role == "protocol" else "per-protocol.")
    )
    doc.add_paragraph(
        f"A total of {120 if role == 'protocol' else 100} participants will be randomised."
    )
    doc.add_paragraph(
        "Primary analysis: analysis of covariance, adjusted for baseline fatigue score."
    )
    doc.add_paragraph(
        "Missing observations will be addressed by multiple imputation. Outcomes after treatment discontinuation will continue to be collected."
    )
    table = doc.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Assessment", "Baseline", "Week 12"]):
        cell.text = value
    for cell, value in zip(table.add_row().cells, ["Fatigue score", "Yes", "Yes"]):
        cell.text = value
    doc.add_paragraph(
        "The trial coordinator arranges primary outcome assessments. Laboratory procedures are delegated to the laboratory manual."
    )
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


class Command(BaseCommand):
    help = "Create a synthetic demonstration project and a local-only account."

    def handle(self, *args, **options):
        if not settings.DEBUG:
            raise CommandError("Demo seeding is restricted to local development (DEBUG=1).")
        user, created = get_user_model().objects.get_or_create(username="demo")
        if created:
            password = get_random_string(20)
            user.set_password(password)
            user.save()
            self.stdout.write(f"Local demo login: demo\nInitial password: {password}")
        else:
            self.stdout.write("Demo account already exists; its password was not changed.")
        project, _ = Project.objects.get_or_create(name="RESTORE · Synthetic study", owner=user)
        project.members.add(user)
        if project.reviews.exists():
            self.stdout.write("Synthetic review already exists.")
            return
        review = Review.objects.create(
            project=project,
            title="Protocol & SAP · Version 1.0",
            scope_confirmed=True,
            metadata_confirmed=True,
        )
        for role in ("protocol", "sap"):
            content = demo_doc(role)
            Document.objects.create(
                review=review,
                role=role,
                name=f"RESTORE-{role}-v1.docx",
                file=ContentFile(content, name=f"{review.id}-{role}.docx"),
                sha256=hashlib.sha256(content).hexdigest(),
            )
        self.stdout.write("Synthetic review queued. Run: python manage.py review_worker --once")
