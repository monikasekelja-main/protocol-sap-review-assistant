import io

from django.utils import timezone
from docx import Document as WordDocument
from docx.shared import Inches, Pt, RGBColor


def export_review(review):
    doc = WordDocument()
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Heading 1", "Heading 2", "Caption"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("244E42")
    doc.styles["Caption"].paragraph_format.keep_with_next = True
    section = doc.sections[0]
    section.header.paragraphs[0].text = "PROTOCOL & SAP REVIEW · REVIEW RECORD"
    section.top_margin = section.bottom_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("Protocol and SAP review report", 0)
    doc.add_paragraph(review.title)
    doc.add_paragraph(
        f"Review snapshot: {timezone.now():%Y-%m-%d %H:%M UTC}\nRun: {review.id} · Record revision {review.revision}"
    )
    doc.add_paragraph(
        "Issues and questions for your trial team. This review checks stated information and alignment; it does not establish scientific, statistical, safety or regulatory adequacy."
    )
    doc.add_heading("Supplied document versions", 1)
    for source in review.documents.all():
        doc.add_paragraph(
            f"{source.get_role_display()}: {source.name}\nTitle: {source.metadata.get('title', 'Unknown')}\nVersion: {source.metadata.get('version', 'Unknown')} · Date: {source.metadata.get('date', 'Unknown')}\nSHA-256: {source.sha256}"
        )

    def evidence(items):
        for item in items:
            doc.add_paragraph(
                f"{item['name']} · version {item['version']} · {item['location']}", style="Caption"
            )
            doc.add_paragraph(item["quote"], style="Quote")

    doc.add_heading("Study summary", 1)
    for field in review.summary.values():
        doc.add_heading(field["field"], 2)
        doc.add_paragraph(field["text"])
        evidence(field["evidence"])
    doc.add_heading("Issues and questions for your trial team", 1)
    for finding in review.findings.all():
        doc.add_heading(finding.title, 2)
        doc.add_paragraph(
            f"{finding.get_kind_display()} · {finding.get_priority_display()} priority (provisional)\nStatus: {finding.get_status_display()}"
        )
        doc.add_paragraph(f"Why it matters: {finding.why}")
        if finding.unknown:
            doc.add_paragraph(f"Unknown: {finding.unknown}")
        doc.add_paragraph(f"Suggested action: {finding.action}")
        evidence(finding.evidence)
        doc.add_paragraph(f"Team note: {finding.note or 'None recorded'}")
    if not review.findings.exists():
        doc.add_paragraph(
            "No supported findings were produced. This is not a pass result; see coverage limitations."
        )
    doc.add_heading("Accepted user clarifications", 1)
    for event in review.events.filter(kind="clarification").select_related("actor"):
        doc.add_paragraph(
            f"{event.created_at:%Y-%m-%d %H:%M UTC} · {event.actor.get_username()}: {event.payload['note']}"
        )
    doc.add_heading("Coverage and limitations", 1)
    doc.add_paragraph(
        "Supported scope: English-language, two-arm parallel randomised superiority trial; continuous primary endpoint at a fixed time point. Source quotations are extracted text; source documents remain unchanged."
    )
    for entry in review.coverage:
        doc.add_paragraph(
            f"{entry['label']}: {entry['state'].replace('_', ' ')} — {entry['reason']}"
        )
    for limitation in review.limitations:
        doc.add_paragraph(limitation, style="List Bullet")
    doc.add_heading("Processing versions", 2)
    for key, value in review.versions.items():
        doc.add_paragraph(f"{key}: {value}")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
